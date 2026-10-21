"""Extract text from supported user documents and prepare vector chunks."""
import csv
import hashlib
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup

from app.retrieval.chunking import chunk_text
from app.retrieval.vector_store import upsert_chunks
from app.services.industry_resolution import resolve_industry


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".html", ".htm", ".csv", ".md", ".markdown"
}


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type: {suffix or '(no extension)'}"
        )

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        paragraphs = [
            p.text for p in document.paragraphs if p.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )
        return "\n".join(paragraphs)

    decoded = content.decode("utf-8-sig", errors="replace")
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(decoded, "html.parser")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        return soup.get_text(" ", strip=True)

    if suffix == ".csv":
        rows = csv.reader(io.StringIO(decoded))
        return "\n".join(
            " | ".join(cell.strip() for cell in row) for row in rows
        )

    return decoded


def ingest_document(
    filename: str,
    content: bytes,
    industry: str,
    domain: str = "DOCUMENT",
    source_url: str | None = None,
    chunk_size: int = 2000,
    overlap: int = 300,
) -> dict[str, Any]:
    resolved = resolve_industry(industry, allow_unknown=True)
    assert resolved is not None

    extracted = extract_text(filename, content)
    pieces = chunk_text(extracted, chunk_size=chunk_size, overlap=overlap)
    if not pieces:
        raise ValueError("No extractable text was found in the document.")

    digest = hashlib.sha256(content).hexdigest()
    source_id = str(uuid.uuid5(
        uuid.NAMESPACE_URL,
        f"{filename}:{digest}:{resolved.id}",
    ))
    now = datetime.now(timezone.utc).isoformat()
    normalised_domain = (domain or "DOCUMENT").strip().upper()

    chunks = []
    for index, piece in enumerate(pieces):
        point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source_id}:{index}"))
        chunks.append({
            "point_id": point_id,
            "text": piece,
            "record_type": "uploaded_document",
            "source_id": source_id,
            "source_name": filename,
            "source_url": source_url or "",
            "industry": resolved.canonical_name,
            "industry_id": resolved.id,
            "industry_ids": [resolved.id],
            # Compatibility field for existing UI and legacy search filters.
            "industries": list(dict.fromkeys([
                resolved.canonical_name, *resolved.aliases
            ])),
            "domain": normalised_domain,
            "chunk_index": index,
            "chunk_count": len(pieces),
            "content_sha256": digest,
            "ingested_at": now,
            "tags": ["uploaded_document"],
            "verified": False,
        })

    upserted = upsert_chunks(chunks)
    return {
        "source_id": source_id,
        "filename": filename,
        "industry": resolved.canonical_name,
        "industry_id": resolved.id,
        "chunks_created_or_updated": upserted,
        "characters_extracted": len(extracted),
    }
